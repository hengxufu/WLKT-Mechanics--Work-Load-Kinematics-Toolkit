from pathlib import Path
import json,sys
from pptx import Presentation
W=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(W/'src/toolchain'))
from receive_web_plan import validate
p=W/'planning/deck_plan.json';plan=json.loads(p.read_text(encoding='utf8'))
for slide,actual in zip(plan['slides'],Presentation(W/'build/deck-native.pptx').slides):
    blocks=[]
    for sh in (list(actual.shapes) if slide['slide_number']==4 else list(actual.shapes)[14:]):
        if sh.name.endswith(('/title','/sources')):continue
        b={'object_id':sh.name,'box_inches':[round(v/914400,5) for v in (sh.left,sh.top,sh.width,sh.height)]}
        if sh.has_text_frame and sh.text.strip():b.update(type='native.text',text=sh.text,source_ref='；'.join(e['source_ref'] for e in slide['evidence']))
        elif sh.has_table:b.update(type='native.table',text=' | '.join(c.text for row in sh.table.rows for c in row.cells),source_ref='sources/live-run-evidence.json')
        elif sh.has_chart:b.update(type='native.chart',data_ref='sources/native-data.json',text='解析预测；非新增实机测量')
        elif sh.shape_type==13:b.update(type='source.image',source_ref=sh.name,text='真实程序截图；对象来源见assets/live与sources/image-dimensions.json')
        else:continue
        blocks.append(b)
    slide['content_blocks']=blocks
validate(plan);p.write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf8')
print('18-page actual geometry synchronized and validated.')
