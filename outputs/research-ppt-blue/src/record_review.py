"""Run only after actually viewing the renders; this records, not performs, visual review."""
from pathlib import Path
import hashlib,json,subprocess,sys
W=Path(__file__).resolve().parents[1];S=W/'src/toolchain'
deck=W/'拉压弯扭大师_蓝色技术路线增强版.pptx'
digest=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
visual=[
'封面左右主次明确，真实工作区与飘带分区，三项能力及定位标签清楚，无遮挡。',
'三类学习问题各有具体疑问、术语与行动，深蓝标题带与浅蓝正文分层，底部闭环可读。',
'四栏等宽，四个透明工程概念素材独立，公式和适用条件保持阅读优先。',
'顶部四类受力示意、五栏流程、校核反馈、当前边界与橙色规划层均为原生对象；缩短长句后无叠字或裁切。',
'修正截图高度后来源行完整露出；截图外侧三项区域解释没有压住真实控件。',
'6m梁、跨中18kN及支座自由度表清楚；荷载属性证据与说明相邻。',
'剪力与弯矩截图等尺度并列，9及负27标注清楚，下侧解释独立。',
'左侧真实证据、右侧原生对照表，放大镜仅占角落；条件、符号、位置说明可读。',
'两张原生图表分开单位，15/12/9及3/6/9反力和15/24/27弯矩正确；解析预测标识可读。',
'四步阶梯及箭头置左，推导截图和解释置右，正文顺序连贯。',
'真实空间截图与概念桁架独立；六自由度标签两行排列，24自由度说明清楚。',
'结果截图、三项结构参数、三向力平衡表分区；残差与未核对力矩说明均完整。',
'计算核心与四模块形成中心构图，修正后箭头从计算核心指向学习解释，芯片不压文字。',
'实现、实验、用途三栏分级清楚；盾牌远离认证文字，边界说明充分。',
'操作指导真实截图与四步说明并列，排错提示独立，字号与间距合适。',
'三个学习价值对应具体已见证据，下一步列表可读，没有空口号代替论证。',
'四格公式手册，符号定义和前提独立，正文可编辑文字公式，无裁切。',
'四类证据采用分行索引，来源及局限独立；末行交流提示完整。']
evidence=[
'软件定位与版本来自README；封面无实验结论。',
'课堂困难为问题归纳；未编教学提升数据。',
'公式为基础线弹性关系；AI金属对象仅概念，不是试验照片或模拟结果。',
'技术路线来自软件说明及能力审计，校核环节独立于软件求解。',
'复用2026-10-10真实程序截图，图标字体异常保留，未宣称本轮重新运行。',
'跨度、荷载、位置及约束与实机记录一致。',
'真实剪力与弯矩证据保留软件负号约定，不篡改结果。',
'仅默认算例幅值对照；9/9kN、27kN·m有解析与实机来源。',
'图表为解析预测，数据源与嵌入工作簿保留，非新增实机测量。',
'推导来自真实软件截图与静力学关系，没有虚构应力结果。',
'24为约束处理前总自由度；概念桁架不冒充实机模型。',
'8000/−5000/−12000N与反力逐向抵消；未宣称完整力矩平衡。',
'Worker CPU及主线程回退有能力审计；WebGPU仅检测。',
'Tet4标为实验、软件标为学习用途；盾牌不表示认证。',
'流程来自真实指导界面与默认案例，不编运行时间。',
'价值对应已展示功能与算例；改进方向表述为未来工作。',
'公式有单位及适用前提，LaTeX源另存，未宣称OMML方程。',
'实际观察、解析预测、实验缺项、视觉协作明确区分；模型名未核验。']
plan=json.loads((W/'planning/deck_plan.json').read_text(encoding='utf8'))
review={'schema':'research_ppt_delivery_review_v1','pptx_sha256':digest(deck),
'web_planning':{'session_url':'https://chatgpt.com/c/6acb57b2-2558-83ee-b982-ec6be3938863','response_file':'planning/visual-assistant-reply.txt','second_review_file':'planning/web-density-review.txt'},
'model_exception':{'instruction':'允许用当前网页版完成规划和单元素生成，注明模型名未核验','observed':'High / 创建图像','verified_model':None},
'review_method':'Actually viewed all 18 final blue renders, including corrected native route slide 4.',
'slides':[],'unresolved_issues':[],'limitations':['Strict GPT-6 selector evidence unavailable; user authorized current web model. Current web follow-up timed out after submission; previous verified planning and 10 accepted visual assets reused.','Real screenshots are reused from 2026-10-10; icon font boxes retained.','Formulas are editable text, not structured Office equations.']}
for n,(slide,v,e) in enumerate(zip(plan['slides'],visual,evidence),1):
    f=W/f'renders/final/幻灯片{n}.PNG'
    review['slides'].append({'slide_number':n,'render_file':str(f.relative_to(W)),'render_sha256':digest(f),'rendered_from_sha256':digest(deck),
      'content_review':{'status':'pass','note':slide['takeaway']+'；已核对具体解释及来源，非仅标题。'},
      'visual_review':{'status':'pass','note':v},'evidence_review':{'status':'pass','note':e}})
(W/'qa/delivery_review.json').write_text(json.dumps(review,ensure_ascii=False,indent=2),encoding='utf8')
from pptx import Presentation
pictures={sh.name:sh for sl in Presentation(deck).slides for sh in sl.shapes if hasattr(sh,'image')}
from datetime import datetime,timezone
for i in range(1,11):
    f=W/f'asset-jobs/R{i:02}/job.json';j=json.loads(f.read_text(encoding='utf8'))
    name=j['spec']['slide_id']+'/'+j['spec']['object_id']
    assert j['status']=='integrated' and digest(W/f'asset-jobs/R{i:02}'/j['asset_file'])==j['sha256']
    assert hashlib.sha256(pictures[name].image.blob).hexdigest()==j['sha256']
    at=datetime.now(timezone.utc).isoformat()
    j['integration']={'pptx':str(deck.resolve()),'pptx_sha256':digest(deck),'shape_name':name,'at':at}
    j['history'].append({'at':at,'status':'integrated','note':'Reused previously accepted asset; actual named picture SHA256 verified in new blue deck; final rendered placement reviewed.'})
    f.write_text(json.dumps(j,ensure_ascii=False,indent=2),encoding='utf8')
print('18 actual page reviews recorded; 10 reused asset hashes verified in blue deck.')
