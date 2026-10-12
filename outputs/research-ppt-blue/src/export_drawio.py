from pathlib import Path
from pptx import Presentation
import xml.etree.ElementTree as E,json,hashlib
W=Path.cwd();prs=Presentation(W/'拉压弯扭大师_技术路线图_可编辑.pptx');ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
r=E.Element('mxfile',host='app.diagrams.net');d=E.SubElement(r,'diagram',name='技术路线',id='route');g=E.SubElement(d,'mxGraphModel',dx='1600',dy='900',grid='1',gridSize='10',page='1',pageWidth='1600',pageHeight='900');root=E.SubElement(g,'root');E.SubElement(root,'mxCell',id='0');E.SubElement(root,'mxCell',id='1',parent='0')
for i,sh in enumerate(prs.slides[0].shapes,2):
    x,y,w,h=[round(v/914400*120,3) for v in (sh.left,sh.top,sh.width,sh.height)];el=sh._element
    fc=el.find('.//a:solidFill/a:srgbClr',ns);lc=el.find('.//a:ln/a:solidFill/a:srgbClr',ns);fill='#'+fc.get('val') if fc is not None else 'none';line='#'+lc.get('val') if lc is not None else 'none'
    preset=el.find('.//a:prstGeom',ns);pr=preset.get('prst') if preset is not None else ''
    if pr=='line':
        xf=el.find('.//a:xfrm',ns);fh=xf is not None and xf.get('flipH')=='1';fv=xf is not None and xf.get('flipV')=='1';arrow=el.find('.//a:tailEnd',ns)
        style=f'strokeColor={line};strokeWidth=1.5;endArrow='+('classic' if arrow is not None and arrow.get('type') not in ('none',None) else 'none')+';'
        c=E.SubElement(root,'mxCell',id=str(i),parent='1',edge='1',style=style);geo=E.SubElement(c,'mxGeometry',relative='1',attrib={'as':'geometry'});E.SubElement(geo,'mxPoint',x=str(x+w if fh else x),y=str(y+h if fv else y),attrib={'as':'sourcePoint'});E.SubElement(geo,'mxPoint',x=str(x if fh else x+w),y=str(y if fv else y+h),attrib={'as':'targetPoint'});continue
    txt=sh.text if sh.has_text_frame else '';style=f'fillColor={fill};strokeColor={line};whiteSpace=wrap;'
    if txt:
        run=next((run for p in sh.text_frame.paragraphs for run in p.runs),None);color='073E32';sz=16
        if run is not None:
            if run.font.size:sz=round(run.font.size.pt*120/72,2)
            try:color=str(run.font.color.rgb)
            except Exception:pass
        style=f'fillColor=none;strokeColor=none;whiteSpace=wrap;align=left;verticalAlign=middle;spacing=0;fontFamily=Microsoft YaHei;fontSize={sz};fontColor=#{color};fontStyle={1 if run and run.font.bold else 0};'
    elif pr=='ellipse':style+='shape=ellipse;'
    elif pr=='triangle':style+='shape=triangle;direction=north;'
    elif pr=='diamond':style+='shape=rhombus;'
    elif pr=='roundRect':style+='rounded=1;arcSize=12;'
    c=E.SubElement(root,'mxCell',id=str(i),parent='1',vertex='1',value=txt,style=style);E.SubElement(c,'mxGeometry',x=str(x),y=str(y),width=str(w),height=str(h),attrib={'as':'geometry'})
f=W/'拉压弯扭大师_技术路线图.drawio';E.indent(r);E.ElementTree(r).write(f,encoding='utf-8',xml_declaration=True);E.parse(f)
(W/'qa/editing-demo.json').write_text(json.dumps({'actual_render_review':[1,6,9],'changes':{'slide_1':'Chinese title changed to 拉压弯扭大师（编辑演示）','slide_6':'native support table A -> A*','slide_9':'native reaction chart first value 15 -> 16'},'delivery_data_unchanged':True,'route_drawio_native_cells':len(prs.slides[0].shapes),'drawio_limit':'Shape-level XML export, not a separately rendered visual acceptance; PPT route is visually reviewed.'},ensure_ascii=False,indent=2),encoding='utf8')
print(f)
