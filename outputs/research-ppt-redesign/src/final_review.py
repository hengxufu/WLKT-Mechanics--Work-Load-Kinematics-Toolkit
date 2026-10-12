from pathlib import Path
import hashlib,json,zipfile
from pptx import Presentation
W=Path(__file__).resolve().parents[1]
P=W/'拉压弯扭大师_中国科研分享重制版.pptx'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
notes=[
('封面明确软件与学习目标，无实验结论','蓝白标题与独立杆件概念插画层级清晰，留白充足','概念图明确标注非实测'),
('三个具体学习问题连接建模与核验','左侧问题与右侧梁示意构成对照，无遮挡','教育目标不作为求解精度证据'),
('四种基本变形附关系式与适用提示','四象限分区与局部圆轴插画可辨，公式为可编辑文本','公式来自本地资料，保留LaTeX源码'),
('五步技术路线含核对回路与两类验证','原生节点及箭头可辨，主次流程与底部检查区分明','路线依据软件实现资料，不宣称全工况通过'),
('解释模型树、视图与属性栏分工','实际工作区大图配右侧三项标注，截图小图标缺字如实保留','复用10月10日本地软件实机截图'),
('6m梁18kN跨中荷载与自由度表一致','梁图与右侧约束表清晰，荷载截图独立放大','两节点一单元与单元内集中力位置来自实机记录'),
('剪力跳变和跨中弯矩峰值对应两图','同等大小截图便于比较，图中9与负27可读','截图符号与幅值分别解释'),
('9/9反力和27弯矩幅值对应解析公式','左证据图右数值摘要，符号说明与端弯矩说明清晰','仅默认案例幅值对比，端弯矩零不等于内部峰值零'),
('荷载位置变化使用给定参数解析预测','两个原生图表分别使用kN和kN·m，标签完整','明确为解析预测非新实机测量，a定义为距左支座距离'),
('整体平衡、局部关系与危险位置形成解释链','真实推导截图与三段原生说明分层，底部提示独立','推导受符号和坐标约定约束'),
('每节点六自由度与12阶单元刚度相连','空间截图左置，右侧六标签和刚度说明无重叠','Euler-Bernoulli假设和放大变形边界明确'),
('四节点三构件24自由度及三向力平衡核验','真实结果表与原生平衡表并列，数值单位一致','残差为界面报告；缺少力矩核验，未填造应力安全系数'),
('三层架构与本地执行路径有具体组件','三层纵向关系与右侧部署模块清楚','GPU仅检测未启用数值内核，依据本地实现资料'),
('能力与边界逐项对应','原生表格列宽可读，学习用途说明独立','不将功能存在等同充分验证，软件用途限制保留'),
('四步复现流程包含单参数核验和排错','截图与四步顺序配合，排错提示可读','入口及行为依据实机指导界面'),
('结论区分已见结果、下一步验证与反馈','大标题、三列内容及独立蓝色环形成收束','未来工作不当作已完成验证'),
('公式与单位及条件逐行对应','原生表格内容完整，最后一行名称自然换行','公式为可编辑文本非OMML，LaTeX源码可追溯'),
('四类证据来源和验证边界完整','来源列表层级清楚，底部边界提示可读','明确网页模型未核验且用户授权继续')]
review={'schema':'research_ppt_delivery_review_v1','pptx_sha256':sha(P),'web_planning':{'session_url':'https://chatgpt.com/c/6acb21ae-4768-83ee-be28-8aa97bceb01e','response_file':'planning/assistant-reply.txt'},'slides':[],'unresolved_issues':[],'accepted_limitations':['用户明确允许当前网页版完成规划和单元素生成，具体模型名称未核验；严格GPT-6型号门禁仍为draft。','公式使用可编辑文本，不是Office原生公式对象。','真实软件截图保留了当前软件部分图标显示缺字现象。'],'no_generated_assets_reason':None,'layout_repetition_reason':None}
for i,ns in enumerate(notes,1):
 r=W/f'renders/final/幻灯片{i}.PNG'
 review['slides'].append({'slide_number':i,'render_file':str(r.relative_to(W)),'render_sha256':sha(r),'rendered_from_sha256':sha(P),**{k:{'status':'pass','note':v} for k,v in zip(['content_review','visual_review','evidence_review'],ns)}})
(W/'qa/delivery_review.json').write_text(json.dumps(review,ensure_ascii=False,indent=2),encoding='utf8')
prs=Presentation(P)
assert len(prs.slides)==18
assert all(s.has_notes_slide for s in prs.slides)
with zipfile.ZipFile(P) as z:
 assert z.testzip() is None
 assert len([n for n in z.namelist() if n.startswith('ppt/charts/chart') and n.endswith('.xml')])==2
demo=Presentation(W/'build/editing-demo-native.pptx')
assert any('编辑演示' in s.text for s in demo.slides[0].shapes if s.has_text_frame)
assert any('编辑演示' in c.text for s in demo.slides[13].shapes if s.has_table for row in s.table.rows for c in row.cells)
chart=next(s.chart for s in demo.slides[8].shapes if s.has_chart)
assert chart.series[0].values[0]==16
(W/'qa/editing-demo.json').write_text(json.dumps({'passed_structural':True,'changes':['第一页标题增加编辑演示','第9页第一图表首值15改16','第14页二维梁表格名称修改'],'source_unchanged':True,'note':'验证副本仅测试对象编辑与渲染，数据修改不用于汇报。'},ensure_ascii=False,indent=2),encoding='utf8')
print('18 page reviews, PPTX integrity, notes and three editing checks recorded.')
